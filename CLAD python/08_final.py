import os
import re
import numpy as np
import pandas as pd
from scipy import stats
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from clad import clad, objective, sensor_clip

# =====================================================
# BAGIAN YANG PERLU KAMU UBAH
# =====================================================
NAMA_FILE  = "Book1.xlsx"
NAMA_SHEET = 0
Y_COL      = "Y"
X_COLS     = ["X1", "X2", "Z", "X2_Z"]     # urutan harus sama dengan sebelumnya
BATAS_KIRI  = 0
BATAS_KANAN = "ARA"
DESIMAL = 4                                 # jumlah desimal di tabel

# Nama yang tampil di tabel dan grafik (boleh diubah, mis. "Y": "Underpricing")
LABEL = {"Konstanta": "Konstanta", "Y": "Y", "X1": "X1", "X2": "X2",
         "Z": "Z", "X2_Z": "X2 x Z", "ARA": "ARA"}
# =====================================================

os.makedirs("output", exist_ok=True)
lab = lambda s: LABEL.get(s, s)

# ---------------- Data ----------------
path = "data/" + NAMA_FILE
if NAMA_FILE.lower().endswith((".xlsx", ".xls")):
    df = pd.read_excel(path, sheet_name=NAMA_SHEET)
else:
    df = pd.read_csv(path, sep=None, engine="python")
kolom = [Y_COL] + X_COLS + [c for c in [BATAS_KIRI, BATAS_KANAN] if isinstance(c, str)]
df = df[kolom].dropna().reset_index(drop=True)

def ambil_batas(spec, n):
    if spec is None:
        return None
    if isinstance(spec, str):
        return df[spec].astype(float).to_numpy()
    return np.full(n, float(spec))

n = len(df)
y = df[Y_COL].astype(float).to_numpy()
X = np.column_stack([np.ones(n)] + [df[c].astype(float).to_numpy() for c in X_COLS])
nama = ["Konstanta"] + X_COLS
L, U = ambil_batas(BATAS_KIRI, n), ambil_batas(BATAS_KANAN, n)
k = X.shape[1]

B = np.load("output/bootstrap_B_final.npy")
b_hat = np.load("output/b_hat.npy")
R = B.shape[0]
if np.isnan(B).any():
    raise SystemExit("Ada NaN di bootstrap_B_final.npy. Jalankan 07_bootstrap_final.py sampai selesai.")

# ---------------- Statistik inti ----------------
se = B.std(axis=0, ddof=1)
z = b_hat / se
p = 2 * (1 - stats.norm.cdf(np.abs(z)))
ci_lo, ci_hi = np.percentile(B, 2.5, axis=0), np.percentile(B, 97.5, axis=0)
V = np.cov(B, rowvar=False, ddof=1)

kiri = int((y <= L + 1e-12).sum()) if L is not None else 0
kanan = int((y >= U - 1e-12).sum()) if U is not None else 0
tidak = n - kiri - kanan

q_model = objective(b_hat, y, X, L, U)
q_const = clad(y, np.ones((n, 1)), L, U, n_acak=20)["q"]
pseudo_r2 = 1 - q_model / q_const

def wald(ix):
    bs = b_hat[ix]
    w = float(bs @ np.linalg.solve(V[np.ix_(ix, ix)], bs))
    return w, len(ix), 1 - stats.chi2.cdf(w, len(ix))

w_all, df_all, p_all = wald([1, 2, 3, 4])

# ---------------- Efek marginal ----------------
x2, zz = X[:, 2], X[:, 3]

def ringkas(est, boot):
    s = boot.std(ddof=1)
    zz_ = est / s
    return [est, s, 2 * (1 - stats.norm.cdf(abs(zz_))),
            np.percentile(boot, 2.5), np.percentile(boot, 97.5)]

me_rows, me_names = [], []
for nm, z0 in [("Z = mean - 1 SD", zz.mean() - zz.std()), ("Z = mean", zz.mean()),
               ("Z = mean + 1 SD", zz.mean() + zz.std())]:
    me_rows.append(ringkas(b_hat[2] + b_hat[4] * z0, B[:, 2] + B[:, 4] * z0))
    me_names.append(f"Efek laten X2 pada {nm} ({z0:.2f})")
for nm, x0 in [("X2 = mean - 1 SD", x2.mean() - x2.std()), ("X2 = mean", x2.mean()),
               ("X2 = mean + 1 SD", x2.mean() + x2.std())]:
    me_rows.append(ringkas(b_hat[3] + b_hat[4] * x0, B[:, 3] + B[:, 4] * x0))
    me_names.append(f"Efek laten Z pada {nm} ({x0:.2f})")
me_rows.append(ringkas(b_hat[1], B[:, 1])); me_names.append("Efek laten X1 (0 ke 1)")

def ame(theta):
    xb = X @ theta
    ind = np.ones(n, dtype=bool)
    if L is not None:
        ind &= xb > L
    if U is not None:
        ind &= xb < U
    ind = ind.astype(float)
    return np.array([ind.mean() * theta[1],
                     np.mean(ind * (theta[2] + theta[4] * zz)),
                     np.mean(ind * (theta[3] + theta[4] * x2)),
                     ind.mean()])

a_hat = ame(b_hat)
a_boot = np.array([ame(b) for b in B])
for j, nm in enumerate(["X1", "X2", "Z"]):
    me_rows.append(ringkas(a_hat[j], a_boot[:, j]))
    me_names.append(f"AME {nm} pada median bersyarat Y teramati")

# ---------------- Format angka ----------------
def f(x, d=None):
    d = DESIMAL if d is None else d
    s = f"{x:.{d}f}"
    return s.lstrip("-") if float(s) == 0 else s

def bintang(pv):
    return "***" if pv < 0.01 else "**" if pv < 0.05 else "*" if pv < 0.10 else ""

deskripsi_batas = (f"batas kiri = {BATAS_KIRI if BATAS_KIRI is not None else 'tidak ada'}, "
                   f"batas kanan = {BATAS_KANAN if BATAS_KANAN is not None else 'tidak ada'}")
sensor_note = "Observasi dianggap tersensor jika Y sama persis dengan batas."
boot_note = f"Standard error dan confidence interval diperoleh dari pairs bootstrap ({R} replikasi)."

# ---------------- Susun tabel ----------------
tabel = []

# Tabel 1: deskriptif
kol_desk = [Y_COL] + X_COLS + ([BATAS_KANAN] if isinstance(BATAS_KANAN, str) else [])
rows = []
for c in kol_desk:
    s = df[c].astype(float)
    rows.append([lab(c), str(int(s.count())), f(s.mean()), f(s.median()),
                 f(s.std()), f(s.min()), f(s.max())])
tabel.append(dict(
    nama="deskriptif", judul="Statistik deskriptif",
    header=["Variabel", "N", "Mean", "Median", "Std. Dev.", "Min", "Max"], rows=rows,
    catatan=(f"Observasi tersensor kiri = {kiri}, tersensor kanan = {kanan}, "
             f"tidak tersensor = {tidak}. {sensor_note}")))

# Tabel 2: regresi
rows = []
for j, nm in enumerate(nama):
    rows.append([lab(nm), f(b_hat[j]) + bintang(p[j])])
    rows.append(["", "(" + f(se[j]) + ")"])
rows += [["Observasi (N)", str(n)],
         ["Tersensor kiri", str(kiri)],
         ["Tersensor kanan", str(kanan)],
         ["Tidak tersensor", str(tidak)],
         ["Pseudo-R²", f(pseudo_r2)],
         [f"Wald χ² ({df_all}), H0: semua slope = 0", f(w_all, 3)],
         ["p-value Wald", f(p_all, 3)],
         ["Replikasi bootstrap", str(R)]]
tabel.append(dict(
    nama="regresi", judul="Hasil estimasi CLAD", header=["Variabel", "CLAD"], rows=rows,
    catatan=(f"Estimator CLAD (Powell, 1984) dengan sensor ganda ({deskripsi_batas}). "
             f"Standard error bootstrap dalam kurung ({R} replikasi, pairs bootstrap). "
             "Signifikansi berdasarkan z = koefisien/SE bootstrap. "
             "*** p<0.01, ** p<0.05, * p<0.10. "
             "Pseudo-R² = 1 - Q(model)/Q(hanya konstanta), dengan Q jumlah selisih absolut. "
             "CLAD tidak berbasis likelihood sehingga log-likelihood, AIC, dan BIC tidak tersedia. "
             + sensor_note)))

# Tabel 3: CI
rows = [[lab(nm), f(b_hat[j]), f(se[j]), f(z[j], 3), f(p[j], 3), f(ci_lo[j]), f(ci_hi[j])]
        for j, nm in enumerate(nama)]
tabel.append(dict(
    nama="confidence_interval", judul="Koefisien CLAD, standard error, dan confidence interval 95%",
    header=["Variabel", "Koefisien", "SE bootstrap", "z", "p-value", "CI 95% bawah", "CI 95% atas"],
    rows=rows, catatan=boot_note + " CI memakai metode percentile."))

# Tabel 4: efek marginal
rows = [[nm, f(r[0]), f(r[1]), f(r[2], 3), f(r[3]), f(r[4])] for nm, r in zip(me_names, me_rows)]
tabel.append(dict(
    nama="efek_marginal", judul="Efek marginal CLAD",
    header=["Efek", "Estimasi", "SE bootstrap", "p-value", "CI 95% bawah", "CI 95% atas"],
    rows=rows,
    catatan=("Efek laten adalah efek pada variabel laten (y*), dihitung dari koefisien dan interaksi. "
             "AME adalah efek marginal rata-rata pada median bersyarat Y teramati, yaitu efek laten "
             "dikalikan indikator bahwa prediksi berada di antara batas kiri dan kanan "
             f"(proporsi observasi tersebut = {f(a_hat[3], 3)}). " + boot_note))
)

# Tabel 5: Wald
uji = {"Semua slope = 0": [1, 2, 3, 4],
       "X2 dan X2 x Z = 0": [2, 4],
       "Z dan X2 x Z = 0": [3, 4],
       "X2, Z, dan X2 x Z = 0": [2, 3, 4]}
rows = []
for nm, ix in uji.items():
    w, d_, pv = wald(ix)
    rows.append([nm, f(w, 3), str(d_), f(pv, 3)])
tabel.append(dict(
    nama="wald", judul="Uji Wald", header=["Hipotesis nol", "Wald χ²", "df", "p-value"],
    rows=rows, catatan="Statistik Wald memakai matriks kovarians hasil bootstrap. " + boot_note))

# ---------------- Ekspor Excel ----------------
with pd.ExcelWriter("output/tabel_final.xlsx", engine="openpyxl") as xw:
    for i, t in enumerate(tabel, 1):
        d = pd.DataFrame(t["rows"], columns=t["header"])
        d.to_excel(xw, sheet_name=f"Tabel{i}_{t['nama']}"[:31], index=False, startrow=1)
        xw.sheets[f"Tabel{i}_{t['nama']}"[:31]].cell(row=1, column=1, value=f"Tabel {i}. {t['judul']}")
        ws = xw.sheets[f"Tabel{i}_{t['nama']}"[:31]]
        ws.cell(row=len(d) + 4, column=1, value="Catatan: " + t["catatan"])
        ws.column_dimensions["A"].width = 52
    pd.DataFrame(B, columns=nama).to_excel(xw, sheet_name="Bootstrap_mentah", index=False)

# ---------------- Ekspor Word ----------------
def border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    for edge, sz in kw.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), "000000")
        b.append(el)

def isi(cell, teks, j, bold=False):
    cell.text = ""
    par = cell.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after = Pt(0)
    run = par.add_run(teks)
    run.font.size = Pt(9.5)
    run.bold = bold

doc = Document()
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(11)
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.0)

for i, t in enumerate(tabel, 1):
    par = doc.add_paragraph()
    r = par.add_run(f"Tabel {i}. {t['judul']}")
    r.bold = True
    r.font.size = Pt(10)
    tb = doc.add_table(rows=len(t["rows"]) + 1, cols=len(t["header"]))
    for j, h in enumerate(t["header"]):
        isi(tb.cell(0, j), h, j, bold=True)
    for a, row in enumerate(t["rows"], 1):
        for j, v in enumerate(row):
            isi(tb.cell(a, j), v, j)
    for c in tb.rows[0].cells:
        border(c, top=12, bottom=6)
    for c in tb.rows[-1].cells:
        border(c, bottom=12)
    par = doc.add_paragraph()
    r = par.add_run("Catatan: " + t["catatan"])
    r.font.size = Pt(8.5)
    doc.add_paragraph()
doc.save("output/tabel_final.docx")

# ---------------- Ekspor LaTeX ----------------
def tex(s):
    s = s.replace("\\", "\\textbackslash{}")
    for a, b in [("&", "\\&"), ("%", "\\%"), ("_", "\\_"), ("#", "\\#")]:
        s = s.replace(a, b)
    s = s.replace("χ²", "$\\chi^2$").replace("²", "$^2$").replace("−", "-").replace("<", "$<$")
    s = re.sub(r"\*+", lambda m: "$^{" + m.group(0) + "}$", s)
    return s

out = ["% Tambahkan di preamble: \\usepackage{booktabs}\n"]
for i, t in enumerate(tabel, 1):
    kolom_fmt = "l" + "c" * (len(t["header"]) - 1)
    out.append("\\begin{table}[htbp]\\centering")
    out.append(f"\\caption{{{tex(t['judul'])}}}\\label{{tab:{t['nama']}}}")
    out.append(f"\\begin{{tabular}}{{{kolom_fmt}}}")
    out.append("\\toprule")
    out.append(" & ".join(tex(h) for h in t["header"]) + " \\\\")
    out.append("\\midrule")
    for row in t["rows"]:
        out.append(" & ".join(tex(v) for v in row) + " \\\\")
    out.append("\\bottomrule")
    out.append("\\end{tabular}")
    out.append("\\begin{minipage}{\\linewidth}\\footnotesize\\vspace{3pt}")
    out.append("\\textit{Catatan:} " + tex(t["catatan"]))
    out.append("\\end{minipage}")
    out.append("\\end{table}\n")
with open("output/tabel_final.tex", "w", encoding="utf-8") as fh:
    fh.write("\n".join(out))

# ---------------- Grafik ----------------
plt.rcParams.update({"font.family": "serif",
                     "font.serif": ["Times New Roman", "DejaVu Serif"],
                     "font.size": 10, "axes.spines.top": False,
                     "axes.spines.right": False, "savefig.dpi": 300})

def simpan(fig, nm):
    fig.savefig(f"output/{nm}.png", bbox_inches="tight")
    fig.savefig(f"output/{nm}.pdf", bbox_inches="tight")
    plt.close(fig)

# Grafik 1: histogram dengan batas sensor
fig, ax = plt.subplots(figsize=(6.2, 4))
ax.hist(y, bins=30, color="0.75", edgecolor="black", linewidth=0.5)
top = ax.get_ylim()[1]
if U is not None:
    for u in np.unique(U):
        ax.axvline(u, color="black", ls="--", lw=1)
        ax.text(u, top * 0.97, f" batas kanan {u:.0%}", rotation=90, va="top", ha="right", fontsize=8)
if L is not None:
    for l_ in np.unique(L):
        ax.axvline(l_, color="black", ls=":", lw=1)
        ax.text(l_, top * 0.97, f" batas kiri {l_:.0%}", rotation=90, va="top", ha="left", fontsize=8)
ax.set_xlabel(lab(Y_COL))
ax.set_ylabel("Frekuensi")
simpan(fig, "grafik1_histogram")

# Grafik 2: aktual vs prediksi
pred = sensor_clip(X @ b_hat, L, U)
cens = np.zeros(n, dtype=bool)
if U is not None:
    cens |= y >= U - 1e-12
if L is not None:
    cens |= y <= L + 1e-12
fig, ax = plt.subplots(figsize=(5, 5))
ax.scatter(pred[~cens], y[~cens], s=18, c="0.3", edgecolors="none", label="Tidak tersensor")
ax.scatter(pred[cens], y[cens], s=26, facecolors="none", edgecolors="black", label="Tersensor")
lim = [min(y.min(), pred.min()), max(y.max(), pred.max())]
ax.plot(lim, lim, "k--", lw=0.8)
ax.set_xlabel("Nilai prediksi CLAD (setelah sensor)")
ax.set_ylabel("Nilai aktual")
ax.legend(frameon=False)
simpan(fig, "grafik2_aktual_vs_prediksi")

# Grafik 3: koefisien terstandar dengan CI
sd = X.std(axis=0, ddof=1)
ix = [1, 2, 3, 4]
est = b_hat[ix] * sd[ix]
lo, hi = ci_lo[ix] * sd[ix], ci_hi[ix] * sd[ix]
ypos = np.arange(len(ix))[::-1]
fig, ax = plt.subplots(figsize=(6, 3.5))
ax.errorbar(est, ypos, xerr=[np.maximum(est - lo, 0), np.maximum(hi - est, 0)],
            fmt="o", color="black", capsize=3)
ax.axvline(0, color="black", lw=0.8)
ax.set_yticks(ypos)
ax.set_yticklabels([lab(nama[i]) for i in ix])
ax.set_xlabel("Koefisien × SD regresor (CI 95% percentile bootstrap)")
simpan(fig, "grafik3_plot_koefisien")

# Grafik 4: distribusi bootstrap
fig, axs = plt.subplots(2, 3, figsize=(9, 5.5))
axs = axs.ravel()
for j in range(k):
    lo_, hi_ = np.percentile(B[:, j], [0.5, 99.5])
    axs[j].hist(B[:, j], bins=40, range=(lo_, hi_), color="0.75", edgecolor="black", linewidth=0.4)
    axs[j].axvline(b_hat[j], color="black", lw=1.2)
    axs[j].axvline(np.median(B[:, j]), color="black", ls="--", lw=1)
    axs[j].set_title(lab(nama[j]), fontsize=10)
for j in range(k, len(axs)):
    axs[j].axis("off")
fig.text(0.5, -0.01, "Garis tegas = estimasi utama; garis putus-putus = median bootstrap. "
         "Sumbu menampilkan persentil 0.5 sampai 99.5.", ha="center", fontsize=8)
fig.tight_layout()
simpan(fig, "grafik4_distribusi_bootstrap")

# Grafik 5: stabilitas SE terhadap jumlah replikasi
ms = np.arange(50, R + 1, 10)
gaya = ["-", "--", "-.", ":", (0, (5, 1))]
fig, ax = plt.subplots(figsize=(6.2, 4))
for j in range(k):
    ratio = np.array([B[:m, j].std(ddof=1) for m in ms]) / se[j]
    ax.plot(ms, ratio, color="black", ls=gaya[j % len(gaya)], lw=1, label=lab(nama[j]))
ax.axhline(1, color="0.5", lw=0.6)
ax.set_xlabel("Jumlah replikasi bootstrap")
ax.set_ylabel("SE kumulatif / SE final")
ax.legend(frameon=False, ncol=2, fontsize=8)
simpan(fig, "grafik5_stabilitas_SE")

# ---------------- Ringkasan di layar ----------------
pd.set_option("display.width", 220)
pd.set_option("display.max_columns", None)
for i, t in enumerate(tabel, 1):
    print("=" * 70)
    print(f"Tabel {i}. {t['judul']}")
    print(pd.DataFrame(t["rows"], columns=t["header"]).to_string(index=False))
print("\nSelesai. File di folder output:")
for fl in ["tabel_final.docx", "tabel_final.xlsx", "tabel_final.tex",
           "grafik1_histogram", "grafik2_aktual_vs_prediksi", "grafik3_plot_koefisien",
           "grafik4_distribusi_bootstrap", "grafik5_stabilitas_SE"]:
    print(" -", fl, "" if "." in fl else "(.png dan .pdf)")